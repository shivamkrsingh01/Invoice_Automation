import io
import pandas as pd
from docx import Document
from typing import Dict, Any, Optional
from app.utils.logging_config import logger


class DocumentProcessor:
    """Process various document types for invoice extraction."""
    
    @staticmethod
    def process_word_document(content: bytes, filename: str) -> Dict[str, Any]:
        """Extract text and tables from Word documents."""
        logger.info(f"Processing Word document: {filename}")
        
        try:
            doc = Document(io.BytesIO(content))
            
            # Extract text from paragraphs
            text_content = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content.append(paragraph.text.strip())
            
            # Extract tables
            tables_data = []
            for table in doc.tables:
                table_rows = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    table_rows.append(row_data)
                if table_rows:
                    tables_data.append(table_rows)
            
            result = {
                "text_content": "\n".join(text_content),
                "tables": tables_data,
                "processing_status": "PROCESSED"
            }
            
            logger.info(f"Extracted {len(text_content)} paragraphs and {len(tables_data)} tables from Word document")
            return result
            
        except Exception as e:
            logger.error(f"Failed to process Word document: {str(e)}")
            return {
                "processing_status": "FAILED",
                "error": f"Word document processing failed: {str(e)}"
            }
    
    @staticmethod
    @staticmethod
    def _cell_to_str(value) -> str:
        """Convert one Excel cell value to a plain string. pandas returns
        native numpy/pandas types (int64, float64, NaT/NaN, Timestamp),
        not strings - but every downstream consumer (build_invoices_from_table,
        the label/date/amount regexes) expects plain strings, and calling
        e.g. .strip() on an int crashes outright. This is the single place
        that normalizes a cell before it goes anywhere else."""
        if value is None:
            return ""
        if isinstance(value, float) and pd.isna(value):
            return ""
        try:
            if pd.isna(value):
                return ""
        except (TypeError, ValueError):
            pass
        if isinstance(value, float):
            # Avoid Excel's habit of turning a whole number into "39702.0"
            if value.is_integer():
                return str(int(value))
            return str(value)
        if hasattr(value, "strftime"):  # pandas Timestamp / datetime.date
            return value.strftime("%d.%m.%Y")
        return str(value)

    @staticmethod
    def process_excel_document(content: bytes, filename: str) -> Dict[str, Any]:
        """Extract data from Excel documents."""
        logger.info(f"Processing Excel document: {filename}")
        
        try:
            # Read Excel file
            excel_file = io.BytesIO(content)
            
            # Try to read all sheets
            all_sheets = pd.read_excel(excel_file, sheet_name=None)
            
            sheets_data = {}
            for sheet_name, df in all_sheets.items():
                # Convert DataFrame to list of lists, normalizing every
                # cell to a plain string (see _cell_to_str) - pandas'
                # native int64/float64/NaN/Timestamp values would otherwise
                # crash the string-based parsing everything else expects.
                sheet_data = [
                    [DocumentProcessor._cell_to_str(cell) for cell in row]
                    for row in df.values.tolist()
                ]
                # Add headers
                if df.columns is not None:
                    header_row = [DocumentProcessor._cell_to_str(c) for c in df.columns.tolist()]
                    sheet_data = [header_row] + sheet_data
                sheets_data[sheet_name] = sheet_data
            
            result = {
                "sheets": sheets_data,
                "processing_status": "PROCESSED"
            }
            
            logger.info(f"Extracted {len(sheets_data)} sheets from Excel document")
            return result
            
        except Exception as e:
            logger.error(f"Failed to process Excel document: {str(e)}")
            return {
                "processing_status": "FAILED",
                "error": f"Excel document processing failed: {str(e)}"
            }
    
    @staticmethod
    def process_email_body(body_text: str) -> Dict[str, Any]:
        """Extract invoice information from email body text."""
        logger.info("Processing email body content")
        
        try:
            # Simple text analysis for invoice patterns
            result = {
                "text_content": body_text,
                "processing_status": "PROCESSED"
            }
            
            # Look for common invoice patterns in text
            if body_text:
                # This is a placeholder for more sophisticated extraction
                # The free_extraction_service will handle the actual invoice extraction
                result["text_content"] = body_text
            
            logger.info(f"Processed email body (length: {len(body_text)} chars)")
            return result
            
        except Exception as e:
            logger.error(f"Failed to process email body: {str(e)}")
            return {
                "processing_status": "FAILED",
                "error": f"Email body processing failed: {str(e)}"
            }